"""
backend/text_extraction.py

강의자료(PDF/PPTX) 파일에서 텍스트만 뽑아내는 유틸리티.
추출한 텍스트는 llm_recommender.generate_quiz / generate_study_list 등에
study_content / material_text로 그대로 넘겨서 쓴다.
"""


def extract_text_from_pdf(file_bytes: bytes) -> str:
    from pypdf import PdfReader
    import io

    reader = PdfReader(io.BytesIO(file_bytes))
    pages = [(page.extract_text() or "") for page in reader.pages]
    return "\n".join(pages).strip()


def _iter_shape_texts(shapes):
    """슬라이드의 모든 도형에서 텍스트를 뽑는다. 표와 그룹(묶인 도형) 안까지 재귀적으로 들어간다
    — 이 두 가지를 빼먹으면 실제 강의자료 PPT(비교표, 다이어그램)에서 내용이 통째로 누락된다."""
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    texts = []
    for shape in shapes:
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            texts.extend(_iter_shape_texts(shape.shapes))
            continue

        if shape.has_table:
            for row in shape.table.rows:
                cells_text = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells_text:
                    texts.append(" | ".join(cells_text))
            continue

        if shape.has_text_frame:
            for paragraph in shape.text_frame.paragraphs:
                text = "".join(run.text for run in paragraph.runs)
                if text.strip():
                    texts.append(text)

    return texts


def extract_text_from_pptx(file_bytes: bytes) -> str:
    from pptx import Presentation
    import io

    prs = Presentation(io.BytesIO(file_bytes))
    slide_texts = []
    for slide in prs.slides:
        texts = _iter_shape_texts(slide.shapes)
        if texts:
            slide_texts.append("\n".join(texts))
    return "\n\n".join(slide_texts).strip()


def extract_text(filename: str, file_bytes: bytes) -> str:
    lower = filename.lower()
    if lower.endswith(".pdf"):
        return extract_text_from_pdf(file_bytes)
    if lower.endswith(".pptx"):
        return extract_text_from_pptx(file_bytes)
    raise ValueError("지원하지 않는 파일 형식입니다. PDF 또는 PPTX 파일만 업로드해주세요.")
